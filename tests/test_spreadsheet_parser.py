"""Testes unitários para o parser de planilhas Excel e CSV."""
import io
import pytest
from datetime import datetime, timezone
import openpyxl
from src.infrastructure.spreadsheet_parser import (
    parse_spreadsheet,
    _normalize_header,
    _clean_str,
    _parse_datetime,
    _map_headers,
)


def create_mock_excel(sheet_name: str = "Patrimônio", rows: list[list] = None) -> bytes:
    """Cria um arquivo Excel em memória para testes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    if rows:
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_normalize_header_and_clean_str():
    assert _normalize_header(None) == ""
    assert _normalize_header("") == ""
    assert _normalize_header("  LOCALIZAÇÃO  ") == "localizacao"
    assert _normalize_header("Nº SÉRIE #1") == "noserie1"
    
    assert _clean_str(None) is None
    assert _clean_str("   ") is None
    assert _clean_str(" HP ") == "HP"
    assert _clean_str(12345) == "12345"


def test_parse_datetime():
    assert _parse_datetime(None) is None
    assert _parse_datetime("   ") is None
    
    # Datetime com e sem timezone
    dt_naive = datetime(2023, 5, 10, 14, 30)
    parsed_naive = _parse_datetime(dt_naive)
    assert parsed_naive.tzinfo == timezone.utc
    assert parsed_naive.year == 2023

    dt_aware = datetime(2023, 5, 10, 14, 30, tzinfo=timezone.utc)
    assert _parse_datetime(dt_aware) == dt_aware
    
    # Formatos de string válidos
    dt1 = _parse_datetime("25/12/2022")
    assert dt1 == datetime(2022, 12, 25, 0, 0, tzinfo=timezone.utc)
    
    dt2 = _parse_datetime("2024-01-15")
    assert dt2 == datetime(2024, 1, 15, 0, 0, tzinfo=timezone.utc)
    
    dt3 = _parse_datetime("15-01-2024")
    assert dt3 == datetime(2024, 1, 15, 0, 0, tzinfo=timezone.utc)

    # Formato inválido
    assert _parse_datetime("data_invalida_xyz") is None


def test_map_headers():
    raw_headers = [
        "TIPO", " LOCALIZAÇÃO", "SITUAÇÃO", "DESCRIÇÃO", "Nº PATRIMÔNIO",
        "Nº SÉRIE", "MARCA", "COLUNA_DESCONHECIDA_XYZ", None, "   "
    ]
    mapping = _map_headers(raw_headers)
    assert mapping["equipment_type"] == 0
    assert mapping["location"] == 1
    assert mapping["status"] == 2
    assert mapping["description"] == 3
    assert mapping["patrimony_number"] == 4
    assert mapping["serial_number"] == 5
    assert mapping["brand"] == 6
    assert "notes" not in mapping


def test_parse_spreadsheet_empty_file_raises():
    with pytest.raises(ValueError, match="O arquivo enviado está vazio"):
        parse_spreadsheet(b"", "teste.xlsx")


def test_parse_spreadsheet_unsupported_or_corrupt_file():
    with pytest.raises(ValueError, match="Formato de arquivo não suportado"):
        parse_spreadsheet(b"not a valid content", "arquivo.txt")


def test_parse_excel_valid():
    headers = [
        "TIPO", "LOCALIZAÇÃO", "SITUAÇÃO", "DESCRIÇÃO", "Nº PATRIMÔNIO",
        "Nº SÉRIE", "MARCA", "ÚLTIMA MANUTENÇÃO", "CHAVE WINDOWS",
        "Nº PRODUTO", "HISTÓRICO DE MOVIMENTAÇÕES", "HOSTNAME", "OBSERVAÇÕES"
    ]
    row1 = [
        "DESKTOP", "TI DEPÓSITO", "EM USO", "Dell OptiPlex 7050", "2020001",
        "SN123456", "Dell", datetime(2022, 10, 1), "WIN-KEY-123",
        "PN-999", "10/10/2022: Limpeza interna", "TI-DESK-01", "Observação teste"
    ]
    row_empty_ident = [
        "MONITOR", "SALA 1", "OCIOSO", "LG 24 Polegadas", None,
        None, "LG", "01/05/2021", None,
        None, None, None, None
    ]
    blank_row = [None, None, None, None, None, None, None, None, None, None, None, None, None]

    excel_bytes = create_mock_excel(
        sheet_name="Ativos Patrimônio",
        rows=[headers, row1, row_empty_ident, blank_row]
    )

    parsed = parse_spreadsheet(excel_bytes, "inventario.xlsx")
    assert parsed.filename == "inventario.xlsx"
    assert parsed.detected_sheet == "Ativos Patrimônio"
    assert len(parsed.rows) == 2

    # Verifica primeiro item
    item1 = parsed.rows[0]
    assert item1.equipment_type == "DESKTOP"
    assert item1.location == "TI DEPÓSITO"
    assert item1.status == "EM USO"
    assert item1.description == "Dell OptiPlex 7050"
    assert item1.patrimony_number == "2020001"
    assert item1.serial_number == "SN123456"
    assert item1.brand == "Dell"
    assert item1.hostname == "TI-DESK-01"
    assert item1.windows_key == "WIN-KEY-123"
    assert item1.product_number == "PN-999"
    assert item1.hist_mov == "10/10/2022: Limpeza interna"
    assert item1.notes == "Observação teste"
    assert item1.last_maintenance_at == datetime(2022, 10, 1, 0, 0, tzinfo=timezone.utc)

    # Verifica segundo item com nulos
    item2 = parsed.rows[1]
    assert item2.equipment_type == "MONITOR"
    assert item2.patrimony_number is None
    assert item2.serial_number is None
    assert item2.brand == "LG"
    assert item2.last_maintenance_at == datetime(2021, 5, 1, 0, 0, tzinfo=timezone.utc)


def test_parse_excel_corrupted_raises():
    with pytest.raises(ValueError, match="Não foi possível abrir o arquivo Excel"):
        parse_spreadsheet(b"corrupted zip bytes", "teste.xlsx")


def test_parse_excel_no_sheetnames(monkeypatch):
    excel_bytes = create_mock_excel(sheet_name="Aba", rows=[["TIPO"], ["DESK"]])
    
    class FakeWb:
        sheetnames = []
    
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *args, **kwargs: FakeWb())
    with pytest.raises(ValueError, match="não contém abas"):
        parse_spreadsheet(excel_bytes, "sem_abas.xlsx")


def test_parse_excel_empty_sheet_raises():
    excel_bytes = create_mock_excel(sheet_name="Planilha Vazia", rows=[])
    with pytest.raises(ValueError, match="está vazia"):
        parse_spreadsheet(excel_bytes, "vazio.xlsx")


def test_parse_csv_valid_semicolon_and_comma_and_latin1():
    csv_content = (
        "TIPO;LOCALIZAÇÃO;SITUAÇÃO;DESCRIÇÃO;Nº PATRIMÔNIO;Nº SÉRIE;MARCA;ÚLTIMA MANUTENÇÃO;OBSERVAÇÕES\n"
        "COMPUTADOR;CTM;EM USO;PC HP 6005;201100;BRG123;HP;15/03/2023;Nota CSV\n"
        "   \n"
        " ; ; ; ; ; ; ; ; \n"
        "SWITCH;REDE;OCIOSO;Switch Cisco;;;Cisco;;\n"
    )
    # Test latin1 encode
    parsed = parse_spreadsheet(csv_content.encode("latin-1"), "dados.csv")
    assert len(parsed.rows) == 2
    assert parsed.rows[0].equipment_type == "COMPUTADOR"
    assert parsed.rows[0].serial_number == "BRG123"
    assert parsed.rows[0].notes == "Nota CSV"
    assert parsed.rows[1].equipment_type == "SWITCH"
    assert parsed.rows[1].patrimony_number is None


def test_parse_csv_empty_raises():
    with pytest.raises(ValueError, match="O arquivo CSV está vazio"):
        parse_spreadsheet(b"   \n\n  ", "vazio.csv")


def test_parse_csv_undecodable_bytes():
    from src.infrastructure import spreadsheet_parser
    
    class FakeBytes:
        def decode(self, *args, **kwargs):
            raise UnicodeDecodeError("utf-8", b"", 0, 1, "invalid")

    with pytest.raises(ValueError, match="Não foi possível decodificar o arquivo CSV"):
        spreadsheet_parser._parse_csv(FakeBytes(), "teste.csv")


def test_parse_csv_no_headers():
    from src.infrastructure import spreadsheet_parser
    with pytest.raises(ValueError, match="O arquivo CSV está vazio"):
        spreadsheet_parser._parse_csv(b"\n\n", "teste.csv")
