"""E2E Test Suite - Tier 2: Boundary & Corner Cases.

Covers boundary limits, invalid inputs, edge cases, RBAC restrictions, and error flows:
- F1: Auth Boundaries & Security Rate Limiting (RN-02, RN-05)
- F2: Equipment Field Uniqueness, Null-Clearing & 404s (RN-EQ-15, TASK-QA-04)
- F3: Maintenance & Movement Edge Conditions & 404s (RN-EQ-11, RN-EQ-12)
- F4: Tag Taxonomy Duplicates & Categories (RN-TAG-02)
- F5: Spreadsheet Parsing Edge Cases & Corrupt Inputs (RN-IMP-03, RN-IMP-06)
- F6: Dashboard, Autocomplete & RBAC Restrictions (RN-DASH-03, RN-EQ-16)
"""
import io

import openpyxl
import pytest
from httpx import AsyncClient

# ==============================================================================
# F1: Authentication & Rate Limiting Boundaries
# ==============================================================================

@pytest.mark.anyio
async def test_tier2_auth_wrong_password_generic_error(client: AsyncClient, e2e_users):
    """Testa que erro de senha incorreta retorna 401 com mensagem genérica sem vazar detalhe (RN-05)."""
    response = await client.post(
        "/api/v1/auth/login",
        json={"email": e2e_users["tecnico"].email, "password": "WrongPassword999!"},
    )
    assert response.status_code == 401
    assert "inválid" in response.json()["detail"].lower()


@pytest.mark.anyio
async def test_tier2_auth_nonexistent_email_generic_error(client: AsyncClient):
    """Testa que email não cadastrado retorna o mesmo status 401 e mensagem genérica (RN-05)."""
    response = await client.post(
        "/api/v1/auth/login",
        json={"email": "naoexiste_e2e@empresa.com", "password": "QualquerSenha123!"},
    )
    assert response.status_code == 401
    assert "inválid" in response.json()["detail"].lower()


@pytest.mark.anyio
async def test_tier2_auth_rate_limit_10_attempts_blocked_on_11th(client: AsyncClient, e2e_users):
    """Testa proteção contra força bruta: 10 falhas pelo mesmo IP geram 401, a 11ª resulta em 429 (RN-02)."""
    ip_headers = {"X-Forwarded-For": "198.51.100.77"}
    bad_payload = {"email": e2e_users["tecnico"].email, "password": "WrongPassword!"}

    # Executa 10 tentativas com falha
    for _ in range(10):
        resp = await client.post("/api/v1/auth/login", json=bad_payload, headers=ip_headers)
        assert resp.status_code == 401

    # A 11ª tentativa deve ser bloqueada imediatamente com 429 Too Many Requests
    blocked_resp = await client.post("/api/v1/auth/login", json=bad_payload, headers=ip_headers)
    assert blocked_resp.status_code == 429
    assert "bloqueado" in blocked_resp.json()["detail"].lower()


@pytest.mark.anyio
async def test_tier2_auth_refresh_missing_cookie_rejected(client: AsyncClient):
    """Testa que rota de renovação sem o cookie refresh_token retorna 401."""
    response = await client.post("/api/v1/auth/refresh")
    assert response.status_code == 401
    assert "ausente" in response.json()["detail"].lower()


@pytest.mark.anyio
async def test_tier2_auth_refresh_invalid_token_type(client: AsyncClient, tecnico_tokens):
    """Testa que passar um access_token no lugar do refresh_token é recusado com 401."""
    client.cookies.set("refresh_token", tecnico_tokens["access_token"])
    response = await client.post("/api/v1/auth/refresh")
    assert response.status_code == 401
    assert "inválido" in response.json()["detail"].lower()


@pytest.mark.anyio
async def test_tier2_auth_tampered_token_rejected(client: AsyncClient):
    """Testa envio de token JWT adulterado ou com assinatura inválida."""
    client.cookies.set("access_token", "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.adulterado.assinaturafalsa")
    response = await client.get("/api/v1/auth/me")
    assert response.status_code == 401


# ==============================================================================
# F2: Equipment Field Uniqueness, Null-Clearing & 404s
# ==============================================================================

@pytest.mark.anyio
async def test_tier2_equipment_duplicate_serial_number_rejected(client: AsyncClient, tecnico_tokens):
    """Testa que número de série duplicado é recusado com HTTP 400 (RN-EQ-15)."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    payload = {"serial_number": "SN-DUP-01", "description": "Equip 1", "equipment_type": "Notebook", "location": "TI"}
    first = await client.post("/api/v1/equipments", json=payload)
    assert first.status_code == 201

    duplicate = await client.post("/api/v1/equipments", json=payload)
    assert duplicate.status_code == 400
    assert "cadastrado" in duplicate.json()["detail"].lower()


@pytest.mark.anyio
async def test_tier2_equipment_duplicate_patrimony_number_rejected(client: AsyncClient, tecnico_tokens):
    """Testa que número de patrimônio duplicado é recusado com HTTP 400."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    p1 = {"patrimony_number": "PAT-DUP-01", "description": "Equip A", "equipment_type": "Desktop", "location": "TI"}
    p2 = {"patrimony_number": "PAT-DUP-01", "description": "Equip B", "equipment_type": "Desktop", "location": "RH"}

    assert (await client.post("/api/v1/equipments", json=p1)).status_code == 201
    dup_resp = await client.post("/api/v1/equipments", json=p2)
    assert dup_resp.status_code == 400
    assert "patrimônio" in dup_resp.json()["detail"].lower()


@pytest.mark.anyio
async def test_tier2_equipment_duplicate_hostname_rejected(client: AsyncClient, tecnico_tokens):
    """Testa que hostname duplicado é recusado com HTTP 400."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    p1 = {"hostname": "SRV-HOST-01", "description": "Servidor 1", "equipment_type": "Servidor", "location": "DC"}
    p2 = {"hostname": "SRV-HOST-01", "description": "Servidor 2", "equipment_type": "Servidor", "location": "DC"}

    assert (await client.post("/api/v1/equipments", json=p1)).status_code == 201
    dup_resp = await client.post("/api/v1/equipments", json=p2)
    assert dup_resp.status_code == 400
    assert "hostname" in dup_resp.json()["detail"].lower()


@pytest.mark.anyio
async def test_tier2_equipment_clear_optional_fields_to_null_success(client: AsyncClient, tecnico_tokens):
    """Testa que limpar campos opcionais para None/vazio não gera falso erro de unicidade (TASK-QA-04)."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    initial_payload = {
        "serial_number": "SN-CLEAR-01",
        "patrimony_number": "PAT-CLEAR-01",
        "hostname": "HOST-CLEAR-01",
        "description": "Equipamento Inicial",
        "equipment_type": "Notebook",
        "location": "TI",
    }
    created = await client.post("/api/v1/equipments", json=initial_payload)
    eq_id = created.json()["id"]

    # Atualiza desvinculando patrimony_number e hostname para None
    clear_payload = {
        "serial_number": "SN-CLEAR-01",
        "patrimony_number": None,
        "hostname": None,
        "description": "Equipamento com campos limpos",
    }
    update_resp = await client.put(f"/api/v1/equipments/{eq_id}", json=clear_payload)
    assert update_resp.status_code == 200
    updated = update_resp.json()
    assert updated["patrimony_number"] is None
    assert updated["hostname"] is None


@pytest.mark.anyio
async def test_tier2_equipment_get_nonexistent_id_404(client: AsyncClient, tecnico_tokens):
    """Testa consulta por ID inexistente retornando HTTP 404."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    response = await client.get("/api/v1/equipments/999999")
    assert response.status_code == 404


@pytest.mark.anyio
async def test_tier2_equipment_update_nonexistent_id_404(client: AsyncClient, tecnico_tokens):
    """Testa atualização de equipamento com ID inexistente retornando HTTP 404."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    response = await client.put("/api/v1/equipments/999999", json={"description": "Fantasma"})
    assert response.status_code == 404


# ==============================================================================
# F3: Maintenance & Movement Edge Conditions & 404s
# ==============================================================================

@pytest.mark.anyio
async def test_tier2_maintenance_nonexistent_equipment_404(client: AsyncClient, tecnico_tokens):
    """Testa cadastro de manutenção para equipamento inexistente retornando 404."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    payload = {"maintenance_date": "2026-09-01T00:00:00Z", "description": "Teste"}
    response = await client.post("/api/v1/equipments/999999/maintenances", json=payload)
    assert response.status_code == 404


@pytest.mark.anyio
async def test_tier2_maintenance_update_nonexistent_maintenance_404(client: AsyncClient, tecnico_tokens):
    """Testa edição de manutenção com ID inexistente retornando 404."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    created = await client.post("/api/v1/equipments", json={"description": "Item", "equipment_type": "TI", "location": "TI"})
    eq_id = created.json()["id"]

    response = await client.put(f"/api/v1/equipments/{eq_id}/maintenances/999999", json={"description": "Novo"})
    assert response.status_code == 404


@pytest.mark.anyio
async def test_tier2_maintenance_delete_nonexistent_maintenance_404(client: AsyncClient, tecnico_tokens):
    """Testa exclusão de manutenção inexistente retornando 404."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    created = await client.post("/api/v1/equipments", json={"description": "Item", "equipment_type": "TI", "location": "TI"})
    eq_id = created.json()["id"]

    response = await client.delete(f"/api/v1/equipments/{eq_id}/maintenances/999999")
    assert response.status_code == 404


@pytest.mark.anyio
async def test_tier2_movement_nonexistent_equipment_404(client: AsyncClient, tecnico_tokens):
    """Testa movimentação de equipamento inexistente retornando 404."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    mov_payload = {
        "origin_location": "TI",
        "destination_location": "RH",
        "movement_date": "2026-09-01T00:00:00Z",
    }
    response = await client.post("/api/v1/equipments/999999/movements", json=mov_payload)
    assert response.status_code == 404


@pytest.mark.anyio
async def test_tier2_maintenance_empty_description_rejected(client: AsyncClient, tecnico_tokens):
    """Testa validação de schema: descrição de manutenção vazia retorna 422."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    created = await client.post("/api/v1/equipments", json={"description": "Equip", "equipment_type": "TI", "location": "TI"})
    eq_id = created.json()["id"]

    bad_payload = {"maintenance_date": "2026-09-01T00:00:00Z", "description": ""}
    response = await client.post(f"/api/v1/equipments/{eq_id}/maintenances", json=bad_payload)
    assert response.status_code == 422


@pytest.mark.anyio
async def test_tier2_equipment_manual_independent_update_last_maintenance(client: AsyncClient, tecnico_tokens):
    """Testa que última manutenção pode ser alterada manualmente sem registros no histórico (RN-EQ-11)."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    created = await client.post("/api/v1/equipments", json={"description": "Equip Teste", "equipment_type": "TI", "location": "TI"})
    eq_id = created.json()["id"]

    manual_date = "2025-01-01T10:00:00Z"
    update_resp = await client.put(f"/api/v1/equipments/{eq_id}", json={"last_maintenance_at": manual_date})
    assert update_resp.status_code == 200

    detail = (await client.get(f"/api/v1/equipments/{eq_id}")).json()
    assert "2025-01-01" in detail["last_maintenance_at"]
    assert len(detail["maintenances"]) == 0  # Sem registros adicionados ao histórico


# ==============================================================================
# F4: Tag Taxonomy Boundaries
# ==============================================================================

@pytest.mark.anyio
async def test_tier2_tag_duplicate_name_same_category_rejected(client: AsyncClient, tecnico_tokens):
    """Testa que criar tag com nome duplicado na mesma categoria é recusado com HTTP 400 (RN-TAG-02)."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    payload = {"category": "tipo", "name": "Servidor Rack"}
    assert (await client.post("/api/v1/equipments/tags", json=payload)).status_code == 201

    dup_resp = await client.post("/api/v1/equipments/tags", json=payload)
    assert dup_resp.status_code == 400
    assert "cadastrad" in dup_resp.json()["detail"].lower()


@pytest.mark.anyio
async def test_tier2_tag_same_name_different_category_allowed(client: AsyncClient, tecnico_tokens):
    """Testa que o mesmo nome de tag em categorias distintas é permitido."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    resp1 = await client.post("/api/v1/equipments/tags", json={"category": "tipo", "name": "Operações"})
    resp2 = await client.post("/api/v1/equipments/tags", json={"category": "localizacao", "name": "Operações"})
    assert resp1.status_code == 201
    assert resp2.status_code == 201


@pytest.mark.anyio
async def test_tier2_tag_empty_name_rejected(client: AsyncClient, tecnico_tokens):
    """Testa que tag com nome vazio é rejeitada com HTTP 422."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    resp = await client.post("/api/v1/equipments/tags", json={"category": "tipo", "name": ""})
    assert resp.status_code == 422


@pytest.mark.anyio
async def test_tier2_tag_update_nonexistent_tag_404(client: AsyncClient, tecnico_tokens):
    """Testa atualização de tag inexistente retornando HTTP 404."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    resp = await client.put("/api/v1/equipments/tags/999999", json={"name": "Inexistente"})
    assert resp.status_code == 404


@pytest.mark.anyio
async def test_tier2_tag_rename_to_existing_duplicate_rejected(client: AsyncClient, tecnico_tokens):
    """Testa que renomear tag para um nome já em uso na mesma categoria retorna HTTP 400."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    await client.post("/api/v1/equipments/tags", json={"category": "localizacao", "name": "Setor A"})
    t2 = await client.post("/api/v1/equipments/tags", json={"category": "localizacao", "name": "Setor B"})

    t2_id = t2.json()["id"]
    rename_resp = await client.put(f"/api/v1/equipments/tags/{t2_id}", json={"name": "Setor A"})
    assert rename_resp.status_code == 400
    assert "já existe" in rename_resp.json()["detail"].lower()


@pytest.mark.anyio
async def test_tier2_tag_delete_nonexistent_tag_404(client: AsyncClient, tecnico_tokens):
    """Testa exclusão de tag inexistente retornando HTTP 404."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    resp = await client.delete("/api/v1/equipments/tags/999999")
    assert resp.status_code == 404


# ==============================================================================
# F5: Spreadsheet Parsing Edge Cases & Corrupt Inputs
# ==============================================================================

@pytest.mark.anyio
async def test_tier2_import_empty_file_rejected(client: AsyncClient, tecnico_tokens):
    """Testa envio de arquivo vazio (0 bytes) retornando HTTP 400."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    files = {"file": ("vazio.xlsx", b"", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    resp = await client.post("/api/v1/equipments/import/preview", files=files)
    assert resp.status_code == 400
    assert "vazio" in resp.json()["detail"].lower()


@pytest.mark.anyio
async def test_tier2_import_corrupted_file_rejected(client: AsyncClient, tecnico_tokens):
    """Testa envio de conteúdo binário corrompido com extensão .xlsx retornando HTTP 400."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    files = {"file": ("corrupto.xlsx", b"NOT_A_VALID_ZIP_OR_EXCEL_STREAM", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    resp = await client.post("/api/v1/equipments/import/preview", files=files)
    assert resp.status_code == 400


@pytest.mark.anyio
async def test_tier2_import_empty_filename_rejected(client: AsyncClient, tecnico_tokens, make_csv_bytes):
    """Testa envio de arquivo com nome vazio retornando HTTP 400."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    files = {"file": ("", make_csv_bytes(rows=[["DESKTOP"]], headers=["TIPO"]), "text/csv")}
    resp = await client.post("/api/v1/equipments/import/preview", files=files)
    assert resp.status_code in (400, 422)


@pytest.mark.anyio
async def test_tier2_import_handles_missing_serial_and_patrimony(client: AsyncClient, tecnico_tokens, make_excel_bytes):
    """Testa que linhas sem serial e sem patrimônio são salvas com None sem erro ou ID artificial (RN-IMP-03)."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    headers = ["TIPO", "LOCALIZAÇÃO", "SITUAÇÃO", "DESCRIÇÃO", "Nº SÉRIE", "Nº PATRIMÔNIO"]
    row = ["DESKTOP", "CTM", "EM USO", "Computador Sem Identificadores", None, None]
    file_bytes = make_excel_bytes(rows=[row], headers=headers)

    files = {"file": ("sem_identificadores.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    resp = await client.post("/api/v1/equipments/import/execute", files=files)
    assert resp.status_code == 200
    assert resp.json()["created_count"] == 1


@pytest.mark.anyio
async def test_tier2_import_blank_rows_ignored_gracefully(client: AsyncClient, tecnico_tokens, make_excel_bytes):
    """Testa que linhas totalmente em branco intercaladas na planilha são ignoradas sem falha."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    headers = ["TIPO", "LOCALIZAÇÃO", "SITUAÇÃO", "DESCRIÇÃO"]
    rows = [
        ["NOTEBOOK", "TI", "EM USO", "Item Válido 1"],
        [None, None, None, None],  # Linha vazia
        ["DESKTOP", "RH", "EM USO", "Item Válido 2"],
    ]
    file_bytes = make_excel_bytes(rows=rows, headers=headers)

    files = {"file": ("intercalado.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    resp = await client.post("/api/v1/equipments/import/execute", files=files)
    assert resp.status_code == 200
    assert resp.json()["created_count"] == 2


@pytest.mark.anyio
async def test_tier2_export_empty_database_valid_header(client: AsyncClient, tecnico_tokens):
    """Testa que exportar quando o banco não possui equipamentos gera planilha válida com as 13 colunas."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    resp = await client.get("/api/v1/equipments/export")
    assert resp.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb["Patrimônio"]
    assert ws.max_row == 1
    assert ws.max_column == 13


# ==============================================================================
# F6: Dashboard, Autocomplete & RBAC Restrictions
# ==============================================================================

@pytest.mark.anyio
async def test_tier2_equipment_suggestions_prefix_too_short(client: AsyncClient, tecnico_tokens):
    """Testa que prefixo com menos de 2 caracteres no endpoint de sugestões retorna HTTP 400."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    resp = await client.get("/api/v1/equipments/suggestions?field=patrimony_number&prefix=A")
    assert resp.status_code == 400
    assert "pelo menos 2 caracteres" in resp.json()["detail"].lower()


@pytest.mark.anyio
async def test_tier2_equipment_suggestions_invalid_field(client: AsyncClient, tecnico_tokens):
    """Testa que solicitar sugestões para campo não permitido retorna HTTP 400."""
    client.cookies.set("access_token", tecnico_tokens["access_token"])
    resp = await client.get("/api/v1/equipments/suggestions?field=campo_invalido&prefix=TESTE")
    assert resp.status_code == 400


@pytest.mark.anyio
async def test_tier2_colaborador_rbac_forbidden_equipments_list(client: AsyncClient, colaborador_tokens):
    """Testa que usuário colaborador recebe 403 Forbidden ao tentar listar equipamentos (RN-EQ-16)."""
    client.cookies.set("access_token", colaborador_tokens["access_token"])
    resp = await client.get("/api/v1/equipments")
    assert resp.status_code == 403


@pytest.mark.anyio
async def test_tier2_colaborador_rbac_forbidden_export(client: AsyncClient, colaborador_tokens):
    """Testa que colaborador recebe 403 Forbidden ao tentar exportar planilha (RN-EXP-01)."""
    client.cookies.set("access_token", colaborador_tokens["access_token"])
    resp = await client.get("/api/v1/equipments/export")
    assert resp.status_code == 403


@pytest.mark.anyio
async def test_tier2_colaborador_rbac_forbidden_import_preview(client: AsyncClient, colaborador_tokens, make_csv_bytes):
    """Testa que colaborador recebe 403 Forbidden ao tentar pré-visualizar importação."""
    client.cookies.set("access_token", colaborador_tokens["access_token"])
    files = {"file": ("teste.csv", make_csv_bytes(rows=[["NOTEBOOK"]], headers=["TIPO"]), "text/csv")}
    resp = await client.post("/api/v1/equipments/import/preview", files=files)
    assert resp.status_code == 403


@pytest.mark.anyio
async def test_tier2_colaborador_rbac_forbidden_tag_management(client: AsyncClient, colaborador_tokens):
    """Testa que colaborador recebe 403 Forbidden ao tentar criar nova tag (RN-TAG-07)."""
    client.cookies.set("access_token", colaborador_tokens["access_token"])
    resp = await client.post("/api/v1/equipments/tags", json={"category": "tipo", "name": "Tag Ilegal"})
    assert resp.status_code == 403
