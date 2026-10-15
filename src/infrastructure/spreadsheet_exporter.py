"""Utilitário para geração e exportação de planilhas Excel (.xlsx) de equipamentos."""

import io
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from src.core.security import mask_windows_product_key
from src.domain.models import Equipment


HEADERS = [
    "TIPO",
    "LOCALIZAÇÃO",
    "SITUAÇÃO",
    "DESCRIÇÃO",
    "Nº PATRIMÔNIO",
    "Nº SÉRIE",
    "MARCA",
    "ÚLTIMA MANUTENÇÃO",
    "CHAVE WINDOWS",
    "MODELO",
    "HISTÓRICO DE MANUTENÇÕES",
    "HOSTNAME",
    "OBSERVAÇÕES",
]

COLUMN_WIDTHS = {
    1: 16.0,  # A - TIPO
    2: 20.0,  # B - LOCALIZAÇÃO
    3: 18.0,  # C - SITUAÇÃO
    4: 38.0,  # D - DESCRIÇÃO
    5: 16.0,  # E - Nº PATRIMÔNIO
    6: 18.0,  # F - Nº SÉRIE
    7: 14.0,  # G - MARCA
    8: 18.0,  # H - ÚLTIMA MANUTENÇÃO
    9: 32.0,  # I - CHAVE WINDOWS
    10: 20.0,  # J - MODELO
    11: 45.0,  # K - HISTÓRICO DE MANUTENÇÕES
    12: 16.0,  # L - HOSTNAME
    13: 45.0,  # M - OBSERVAÇÕES
}


def _format_maintenances(equipment: Equipment) -> str:
    """Formata as manutenções cadastradas em ordem cronológica crescente com quebras de linha."""
    if not equipment.maintenances:
        return ""

    sorted_maints = sorted(
        equipment.maintenances,
        key=lambda m: m.started_at or datetime.min.replace(tzinfo=timezone.utc),
    )

    lines: list[str] = []
    for m in sorted_maints:
        date_str = m.started_at.strftime("%d/%m/%Y") if m.started_at else ""
        type_str = (
            m.maintenance_type.name if (m.maintenance_type and m.maintenance_type.name) else None
        )
        desc = m.description.strip() if (m.description and m.description.strip()) else type_str
        notes = f" (Obs: {m.notes.strip()})" if m.notes and m.notes.strip() else ""

        entry = ""
        if date_str and desc:
            entry = f"{date_str}: {desc}{notes}"
        elif desc:
            entry = f"{desc}{notes}"
        elif date_str:
            entry = f"{date_str}{notes}"
        elif notes:
            entry = notes.strip()

        if entry:
            lines.append(entry)

    return "\n".join(lines)


def _format_date(dt: datetime | None) -> str:
    """Formata uma data/hora no padrão DD/MM/YYYY."""
    if not dt:
        return ""
    return dt.strftime("%d/%m/%Y")


def _sanitize_formula_value(val: str | None) -> str:
    """Sanitiza strings para prevenir vulnerabilidade de CSV / Formula Injection (CWE-1236)."""
    if val is None:
        return ""
    val_str = str(val)
    if val_str.startswith(("=", "+", "-", "@")):
        return f"'{val_str}"
    return val_str


def generate_equipments_excel(equipments: list[Equipment]) -> bytes:
    """
    Gera uma planilha Excel (.xlsx) contendo a aba 'Patrimônio' estilizada
    fielmente ao modelo PATRIMÔNIO CTM.xlsx.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Patrimônio"

    # Estilos do cabeçalho
    header_fill = PatternFill(start_color="4DD0E1", end_color="4DD0E1", fill_type="solid")
    header_font = Font(name="Arial", size=9, bold=True, color="000000")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Estilos das células de dados
    data_font = Font(name="Arial", size=9, color="000000")
    data_alignment_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    data_alignment_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    # Inserção do cabeçalho na Linha 1
    ws.row_dimensions[1].height = 26.0
    for col_idx, header_title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    # Inserção dos dados
    for row_idx, eq in enumerate(equipments, start=2):
        ws.row_dimensions[row_idx].height = 20.0

        latest_maint = None
        if eq.maintenances:
            latest_maint = max(
                (m.started_at for m in eq.maintenances if m.started_at), default=None
            )

        type_name = eq.equipment_type.name if eq.equipment_type else ""
        sector_name = eq.sector.name if eq.sector else ""
        status_name = eq.status.name if eq.status else ""

        raw_row_values = [
            type_name,
            sector_name,
            status_name,
            eq.description or "",
            eq.patrimony_number or "",
            eq.serial_number or "",
            eq.brand or "",
            _format_date(latest_maint),
            mask_windows_product_key(eq.windows_product_key_encrypted) or "",
            eq.model or "",
            _format_maintenances(eq),
            eq.hostname or "",
            eq.notes or "",
        ]

        row_values = [_sanitize_formula_value(val) for val in raw_row_values]

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in (1, 3, 5, 6, 7, 8, 10, 12):
                cell.alignment = data_alignment_center
            else:
                cell.alignment = data_alignment_left

    # Configuração de largura das colunas
    for col_idx, width in COLUMN_WIDTHS.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # Congela a linha do cabeçalho
    ws.freeze_panes = "A2"

    # Exibe linhas de grade
    ws.views.sheetView[0].showGridLines = True

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
